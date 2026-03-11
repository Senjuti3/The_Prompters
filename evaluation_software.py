"""
Evaluation Software for Team Image Accuracy Scoring
Processes Excel files from Google Forms and scores image accuracy
"""

import pandas as pd
import numpy as np
from PIL import Image
import requests
from io import BytesIO
import os
import json
from pathlib import Path
from typing import Dict, List, Tuple, Optional
import cv2
from skimage.metrics import structural_similarity as ssim
from openpyxl import load_workbook
import warnings
warnings.filterwarnings('ignore')


class ImageComparator:
    """Handles image comparison and scoring"""
    
    def __init__(self, reference_image_path: str):
        """
        Initialize with reference image path
        
        Args:
            reference_image_path: Path to the reference image file
        """
        self.reference_image_path = reference_image_path
        self.reference_image = None
        self.load_reference_image()
    
    def load_reference_image(self):
        """Load the reference image"""
        if os.path.exists(self.reference_image_path):
            self.reference_image = cv2.imread(self.reference_image_path)
            if self.reference_image is None:
                raise ValueError(f"Could not load reference image from {self.reference_image_path}")
        else:
            raise FileNotFoundError(f"Reference image not found: {self.reference_image_path}")
    
    def download_image_from_url(self, url: str) -> Optional[np.ndarray]:
        """
        Download image from Google Drive URL
        
        Args:
            url: Google Drive shareable link
            
        Returns:
            Image as numpy array or None if download fails
        """
        try:
            # Normalize to string and clean
            if url is None or (isinstance(url, float) and pd.isna(url)):
                return None
            url = str(url).strip()
            if not url:
                return None

            # Convert Google Drive shareable link to direct download link
            if "drive.google.com" in url:
                if "/d/" in url:
                    file_id = url.split("/d/")[1].split("/")[0]
                elif "id=" in url:
                    file_id = url.split("id=")[1].split("&")[0]
                else:
                    file_id = None

                if file_id:
                    direct_url = f"https://drive.google.com/uc?export=download&id={file_id}"
                else:
                    direct_url = url
            else:
                direct_url = url

            response = requests.get(direct_url, timeout=30)
            response.raise_for_status()
            
            img = Image.open(BytesIO(response.content))
            img_array = np.array(img)
            
            # Convert to BGR if RGB
            if len(img_array.shape) == 3 and img_array.shape[2] == 3:
                img_array = cv2.cvtColor(img_array, cv2.COLOR_RGB2BGR)
            elif len(img_array.shape) == 2:
                img_array = cv2.cvtColor(img_array, cv2.COLOR_GRAY2BGR)
            
            return img_array
        except Exception as e:
            print(f"Error downloading image from {url}: {str(e)}")
            return None
    
    def compare_images(self, image_url: str) -> float:
        """
        Compare submitted image with reference image
        
        Args:
            image_url: URL of the submitted image
            
        Returns:
            Similarity score between 0 and 1 (1 = identical, 0 = completely different)
        """
        if image_url is None or pd.isna(image_url) or str(image_url).strip() == '':
            return 0.0
        
        submitted_image = self.download_image_from_url(image_url)
        if submitted_image is None:
            return 0.0
        
        try:
            # Resize images to same dimensions for comparison
            ref_height, ref_width = self.reference_image.shape[:2]
            submitted_image = cv2.resize(submitted_image, (ref_width, ref_height))
            
            # Convert to grayscale for SSIM comparison
            ref_gray = cv2.cvtColor(self.reference_image, cv2.COLOR_BGR2GRAY)
            sub_gray = cv2.cvtColor(submitted_image, cv2.COLOR_BGR2GRAY)
            
            # Calculate Structural Similarity Index
            similarity_score = ssim(ref_gray, sub_gray)
            
            # Normalize to 0-1 range (SSIM already returns -1 to 1, but typically 0 to 1)
            return max(0.0, min(1.0, similarity_score))
        except Exception as e:
            print(f"Error comparing images: {str(e)}")
            return 0.0


class EvaluationProcessor:
    """Main processor for evaluation workflow"""
    
    def __init__(self, config_path: str = "config.json"):
        """
        Initialize processor with configuration
        
        Args:
            config_path: Path to configuration JSON file
        """
        self.config = self.load_config(config_path)
        self.results = []
    
    def load_config(self, config_path: str) -> Dict:
        """Load configuration from JSON file"""
        if os.path.exists(config_path):
            with open(config_path, 'r') as f:
                return json.load(f)
        else:
            # Default configuration
            default_config = {
                "section1_range": [1, 200],
                "section2_range": [201, 400],
                "reference_images": {
                    "question1": "image1.jpg",
                    "question2": "image2.jpg",
                    "question3": "image3.jpg",
                    "question4": "image4.jpg",
                    "question5": "image5.jpg",
                    "question6": "image6.jpg",
                    "question7": "image7.jpg",
                    "question8": "image8.jpg",
                    "question9": "image9.jpg",
                    "question10": "image10.jpg"
                },
                "team_number_column": "Team Number",
                "image_link_column": "Image Link"
            }
            self.save_config(config_path, default_config)
            return default_config
    
    def save_config(self, config_path: str, config: Dict):
        """Save configuration to JSON file"""
        with open(config_path, 'w') as f:
            json.dump(config, f, indent=4)
    
    def process_excel_file(self, excel_path: str, question_number: int, section: int = 1) -> pd.DataFrame:
        """
        Process a single Excel file for one question
        
        Args:
            excel_path: Path to Excel file from Google Form
            question_number: Question number (1-10)
            section: Section number (1 or 2)
            
        Returns:
            DataFrame with team numbers and scores
        """
        print(f"\nProcessing Question {question_number}, Section {section}...")
        print(f"Reading Excel file: {excel_path}")
        
        # Read Excel file
        try:
            df = pd.read_excel(excel_path)
        except Exception as e:
            print(f"Error reading Excel file: {str(e)}")
            return None

        # Keep original Excel row numbers (assuming header on first row)
        df["_ExcelRow"] = df.index + 2
        
        # Find team number and image link columns
        # 1) Try explicit column names from config.json
        team_col = None
        image_col = None

        cfg_team_col = self.config.get("team_number_column")
        cfg_image_col = self.config.get("image_link_column")

        if cfg_team_col and cfg_team_col in df.columns:
            team_col = cfg_team_col
        if cfg_image_col and cfg_image_col in df.columns:
            image_col = cfg_image_col

        # 2) If still missing, try automatic detection (case-insensitive)
        if team_col is None or image_col is None:
            for col in df.columns:
                col_lower = str(col).lower()
                if team_col is None and "team" in col_lower and ("number" in col_lower or "no" in col_lower or "#" in col_lower):
                    team_col = col
                # Treat columns like "Upload image 1/2/..." or "Image link/URL" as image link columns
                if image_col is None:
                    if "image" in col_lower and ("link" in col_lower or "url" in col_lower or "upload" in col_lower):
                        image_col = col

        # 3) Final fallback if still not found
        if team_col is None or image_col is None:
            print(f"Warning: Could not find required columns.")
            print(f"  Available columns: {list(df.columns)}")
            print(f"  Config expects team_number_column='{cfg_team_col}', image_link_column='{cfg_image_col}'")
            print(f"  Falling back to: first column as team number, second as image link")
            team_col = df.columns[0]
            image_col = df.columns[1] if len(df.columns) > 1 else None
        
        if image_col is None:
            print("Error: No image link column found")
            return None
        
        # Extract team numbers, image links and original row index
        df_processed = df[[team_col, image_col, "_ExcelRow"]].copy()
        df_processed.columns = ["Team_Number", "Image_Link", "_ExcelRow"]
        
        # Clean team numbers
        df_processed["Team_Number"] = pd.to_numeric(df_processed["Team_Number"], errors="coerce")
        df_processed = df_processed.dropna(subset=["Team_Number"])
        df_processed["Team_Number"] = df_processed["Team_Number"].astype(int)
        
        # Get reference image path
        question_key = f"question{question_number}"
        reference_image = self.config['reference_images'].get(question_key, f"image{question_number}.jpg")
        
        if not os.path.exists(reference_image):
            print(f"Warning: Reference image '{reference_image}' not found. Please check config.json")
            # Set all scores to 0 if reference image missing
            df_processed["Score"] = 0.0
        else:
            # Prepare Excel workbook to read real hyperlinks (Google Forms often stores "View" as text + real URL as hyperlink)
            try:
                wb = load_workbook(excel_path, data_only=True)
                ws = wb.active

                # Find the column letter for the image link header
                image_col_letter = None
                for cell in ws[1]:
                    if str(cell.value).strip() == str(image_col):
                        image_col_letter = cell.column_letter
                        break
            except Exception as e:
                print(f"Warning: Could not open workbook with openpyxl to read hyperlinks: {e}")
                wb = None
                ws = None
                image_col_letter = None

            # Compare images
            comparator = ImageComparator(reference_image)
            print(f"Comparing images with reference: {reference_image}")

            scores = []
            for _, row in df_processed.iterrows():
                # Start with whatever text pandas read in the cell
                image_url = row["Image_Link"]

                # If we have access to the workbook and column letter, try to get the real hyperlink target
                if ws is not None and image_col_letter is not None:
                    try:
                        excel_row = int(row["_ExcelRow"])
                        cell_ref = f"{image_col_letter}{excel_row}"
                        cell = ws[cell_ref]
                        if cell.hyperlink and cell.hyperlink.target:
                            image_url = cell.hyperlink.target
                    except Exception as e:
                        # If anything fails, just fall back to the cell text
                        print(f"Warning: could not read hyperlink for row {row.get('Team_Number', '?')}: {e}")

                print(f"  Processing Team {row['Team_Number']}...", end="\r")
                score = comparator.compare_images(image_url)
                scores.append(score)

            print()  # New line after progress

            if wb is not None:
                wb.close()

            df_processed["Score"] = scores
        
        # Determine section range
        if section == 1:
            start_team, end_team = self.config['section1_range']
        else:
            start_team, end_team = self.config['section2_range']
        
        # Create complete team list for the section
        all_teams = pd.DataFrame({
            'Team_Number': range(start_team, end_team + 1)
        })
        
        # Merge with submitted data
        result_df = all_teams.merge(
            df_processed[['Team_Number', 'Score']],
            on='Team_Number',
            how='left'
        )
        
        # Fill missing scores with 0
        result_df['Score'] = result_df['Score'].fillna(0.0)
        
        # Sort by team number
        result_df = result_df.sort_values('Team_Number').reset_index(drop=True)
        
        # Rename score column to include question number
        result_df.rename(columns={'Score': f'Q{question_number}_Score'}, inplace=True)
        
        return result_df
    
    def process_all_questions(self, excel_files: List[str], section: int = 1) -> pd.DataFrame:
        """
        Process all 10 questions and combine results
        
        Args:
            excel_files: List of 10 Excel file paths (one per question)
            section: Section number (1 or 2)
            
        Returns:
            Combined DataFrame with all scores
        """
        if len(excel_files) != 10:
            print(f"Warning: Expected 10 Excel files, got {len(excel_files)}")
        
        # Determine section range for team list
        if section == 1:
            start_team, end_team = self.config['section1_range']
        else:
            start_team, end_team = self.config['section2_range']
        
        # Initialize result dataframe with all teams
        result_df = pd.DataFrame({
            'Team_Number': range(start_team, end_team + 1)
        })
        
        # Process each question
        for question_num in range(1, min(len(excel_files) + 1, 11)):
            excel_path = excel_files[question_num - 1]
            
            if not os.path.exists(excel_path):
                print(f"Warning: Excel file not found: {excel_path}")
                # Add column with zeros
                result_df[f'Q{question_num}_Score'] = 0.0
                continue
            
            question_result = self.process_excel_file(excel_path, question_num, section)
            
            if question_result is not None:
                # Merge scores
                result_df = result_df.merge(
                    question_result[['Team_Number', f'Q{question_num}_Score']],
                    on='Team_Number',
                    how='left'
                )
                result_df[f'Q{question_num}_Score'] = result_df[f'Q{question_num}_Score'].fillna(0.0)
            else:
                result_df[f'Q{question_num}_Score'] = 0.0
        
        # Calculate total score
        score_columns = [col for col in result_df.columns if '_Score' in col]
        result_df['Total_Score'] = result_df[score_columns].sum(axis=1)
        
        # Sort by team number
        result_df = result_df.sort_values('Team_Number').reset_index(drop=True)
        
        return result_df
    
    def save_results(self, df: pd.DataFrame, output_path: str):
        """Save results to Excel file"""
        df.to_excel(output_path, index=False)
        print(f"\nResults saved to: {output_path}")


def main():
    """Main function to run the evaluation software"""
    import sys
    
    print("=" * 60)
    print("Evaluation Software - Team Image Accuracy Scoring")
    print("=" * 60)
    
    # Initialize processor
    processor = EvaluationProcessor("config.json")
    
    # Check if running in interactive mode or command line mode
    if len(sys.argv) > 1:
        # Command line mode
        if sys.argv[1] == '--help' or sys.argv[1] == '-h':
            print("\nUsage:")
            print("  python evaluation_software.py")
            print("    Run in interactive mode")
            print("\n  python evaluation_software.py <section> <excel1> <excel2> ... <excel10>")
            print("    Process files directly")
            print("\n  Example:")
            print("    python evaluation_software.py 1 q1.xlsx q2.xlsx ... q10.xlsx")
            return
        
        section = int(sys.argv[1])
        excel_files = sys.argv[2:]
        
        if len(excel_files) < 10:
            print(f"Warning: Only {len(excel_files)} files provided, expected 10")
        
        result_df = processor.process_all_questions(excel_files, section)
        output_file = f"results_section{section}.xlsx"
        processor.save_results(result_df, output_file)
    else:
        # Interactive mode
        print("\nInteractive Mode")
        print("-" * 60)
        
        section = input("Enter section number (1 or 2): ").strip()
        try:
            section = int(section)
            if section not in [1, 2]:
                print("Invalid section number. Using section 1.")
                section = 1
        except:
            print("Invalid input. Using section 1.")
            section = 1
        
        print(f"\nPlease provide paths to 10 Excel files (one per question):")
        excel_files = []
        for i in range(1, 11):
            file_path = input(f"Question {i} Excel file path: ").strip().strip('"')
            excel_files.append(file_path)
        
        result_df = processor.process_all_questions(excel_files, section)
        output_file = input("\nOutput file name (default: results_section{}.xlsx): ".format(section)).strip()
        if not output_file:
            output_file = f"results_section{section}.xlsx"
        
        processor.save_results(result_df, output_file)
        
        print("\n" + "=" * 60)
        print("Processing Complete!")
        print("=" * 60)


if __name__ == "__main__":
    main()
